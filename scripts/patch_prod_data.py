import openpyxl
from datetime import datetime

wb = openpyxl.load_workbook('C:\\Users\\CorySchneider\\AppData\\Local\\CSPM\\Data\\CSPM.xlsm', keep_vba=True)

# 1. Update Transactions sheet row 350
tx_sheet = wb['Transactions']
for row_idx, row in enumerate(tx_sheet.iter_rows(values_only=True), start=1):
    if row[0] == 'TXN-AP-871BEC5F6A1B9C42F3D7':
        tx_sheet.cell(row=row_idx, column=22, value=100) # BillClaimPct
        tx_sheet.cell(row=row_idx, column=23, value=491.06) # TotalClaimAmount
        print(f"Updated TXN row {row_idx}")
        break

# 2. Add row to Disbursements
disb_sheet = wb['Disbursements']
new_row = [
    "DISB-20576109-PATCH", # 0 DisbursementID
    "2026-09-23", # 1 Date
    "Truexperiences Tours Inc.", # 2 ClientName
    None, # 3 SubClient
    "TRU1", # 4 ClientID
    "TRU1", # 5 ParentID
    "M_c598137933", # 6 MatterID
    "CIPO Filing Fees - supplier invoice 20576109", # 7 Description
    491.06, # 8 Amount
    "Y", # 9 TaxExempt
    100, # 10 BillPct
    None, # 11 InvoiceRef
    "PENDING", # 12 PaymentStatus
    None, # 13 InvoiceTotal
    None, # 14 InvoiceAmountPaid
    None, # 15 InvoiceBalanceDue
    None, # 16 ReissueInvoiceNum
    datetime.now().strftime("%Y-%m-%d %H:%M:%S"), # 17 CreatedAt
    "APB-1790203230387", # 18 APBillID
    "APA-PATCH", # 19 APAllocationID
    "TXN-AP-871BEC5F6A1B9C42F3D7", # 20 SourceTransactionID
    "CAD", # 21 OriginalCurrency
    None, # 22 OriginalAmount
    None, # 23 FXRate
    "20576109", # 24 SupplierInvoiceRef
    None, # 25 DocumentPath
]

disb_sheet.append(new_row)
print(f"Appended DISB row: {new_row}")

wb.save('C:\\Users\\CorySchneider\\AppData\\Local\\CSPM\\Data\\CSPM.xlsm')
print("Saved CSPM.xlsm")
