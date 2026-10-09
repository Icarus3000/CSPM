"""Development-only manual adapter; never imported by ordinary CSPM startup."""
from __future__ import annotations

import json
import os
from pathlib import Path
import shutil
import sys

LABEL = "CSPM NATIVE MOTION VISUAL TEST — DISPOSABLE"


def validate_manual_options(args):
    """One collector-free cold configuration; no prepared target or fallback."""
    forbidden = (
        "pixels", "endpoint_pixels", "physical_diagnostics", "source_observation_only",
        "source_unlocked", "source_markers", "controlled_backdrop", "probe_composition",
        "probe_endpoint", "post_input_pixels", "prepared_target", "capture_only",
        "layout_only", "profile_boundaries", "fanout_variant", "input_witness",
        "input_trace", "activation_profile", "qt_render_timings", "qt_layout_polish",
        "profile_python_commit", "source_transfer_trace", "native_call_trace",
    )
    if any(getattr(args, name, False) for name in forbidden):
        raise ValueError("Manual assessment requires no collector, profiling, prepared target or isolation control")
    if (args.duration_ms, args.blend_start_ms, args.gui_delay_ms, args.endpoint_hold_ms) != (350, 240, 0, 0):
        raise ValueError("Manual assessment preserves the existing cold 350/240/100 ms contracts and zero intentional delay")
    if not (args.intrinsic_only and args.single_owner_source and args.native_created_source_band
            and args.render_target_import and args.layout_repair and args.activation_repair):
        raise ValueError("Manual assessment requires the current created-band collector-free configuration")
    if args.workspace != "time-entry" or args.first_direction != "maximize":
        raise ValueError("Manual assessment uses a restored synthetic Time Entry workspace")
    if os.environ.get("CSPM_EXPERIMENTAL_REDUCED_MOTION") == "1":
        raise ValueError("Reduced motion requested; manual native motion launch refused")


def bridge_path(root, filename):
    directory = (Path(root) / "outputs/native_cleanroom").resolve()
    candidate = (directory / filename).resolve()
    if candidate.parent != directory or candidate.suffix.lower() != ".dll" or not candidate.is_file():
        raise ValueError("Manual bridge is missing or outside governed native outputs; no fallback")
    return candidate


def checked_child(root, candidate):
    root, candidate = Path(root).resolve(), Path(candidate).resolve()
    if candidate == root or not candidate.is_relative_to(root):
        raise ValueError("Disposable path must remain inside its task-owned root")
    return candidate


def prepare_profile(root, audit, restored_size):
    """Seed only repository templates; never discover or copy a user profile."""
    from openpyxl import load_workbook
    profile = checked_child(audit, Path(audit) / "disposable_profile")
    local = checked_child(profile, profile / "local")
    local.mkdir(parents=True, exist_ok=False)
    for name in ("CSPM.xlsm", "Dockets.xlsm"):
        shutil.copy2(Path(root) / "src/templates" / name, local / name)
    book = load_workbook(local / "CSPM.xlsm", keep_vba=True)

    def row(sheet, values):
        headers = [cell.value for cell in book[sheet][1]]
        if not set(values).issubset(headers):
            raise ValueError("Synthetic template schema differs from the reviewed checkpoint")
        book[sheet].append([values.get(header) for header in headers])

    for index in range(1, 4):
        client_id, matter_id = f"VISUAL-C{index}", f"VISUAL-M{index}"
        row("Clients", {"ClientID": client_id, "ClientName": f"Synthetic Visual Client {index}",
                        "Status": "Active", "Active": True, "Notes": "DISPOSABLE VISUAL TEST ONLY"})
        row("Matters", {"MatterID": matter_id, "MatterNumber": f"VISUAL-{index:03}",
                        "MatterName": f"Synthetic Matter {index}", "DisplayName": f"Synthetic Matter {index}",
                        "ClientID": client_id, "ClientName": f"Synthetic Visual Client {index}",
                        "Status": "Open", "DefaultRate": 250, "DefaultSharePct": 100})
        row("TimeEntries", {"EntryID": f"VISUAL-E{index}", "Date": "2026-10-01",
                            "ClientID": client_id, "MatterID": matter_id,
                            "Description": f"Synthetic visual fixture entry {index}", "Hours": 1.0,
                            "ClientRate": 250, "SharePct": 100, "GrossToClient": 250,
                            "AmountToYou": 250, "HST": 32.5, "TotalInclHST": 282.5,
                            "RawSeconds": 3600, "Status": "Unbilled"})
    from openpyxl.utils import get_column_letter
    for sheet in ("Clients", "Matters", "TimeEntries"):
        for table in book[sheet].tables.values():
            table.ref = f"A1:{get_column_letter(book[sheet].max_column)}{book[sheet].max_row}"
    book.save(local / "CSPM.xlsm")
    book.close()
    if book.vba_archive:
        book.vba_archive.close()
    settings = {"appStyle": "Professional", "theme": "Dark", "localDataDir": str(local),
                "masterDataDir": "", "keepTrayAlive": False, "runAtStartup": False,
                "soundEffectsEnabled": False, "autoBackupMinutes": 0, "autoBackupIntervalMins": 0,
                "mainWindowLayout": {"maximized": False, "hasExactRect": True,
                    "x": 120, "y": 80, "width": restored_size[0], "height": restored_size[1],
                    "workAreaX": 0, "workAreaY": 0, "workAreaWidth": 1920, "workAreaHeight": 1040}}
    (profile / "user_settings.json").write_text(json.dumps(settings, indent=2), encoding="utf-8")
    overrides = {"LOCALAPPDATA": profile / "os_local", "APPDATA": profile / "os_roaming",
                 "USERPROFILE": profile / "home", "HOME": profile / "home",
                 "TEMP": profile / "tmp", "TMP": profile / "tmp",
                 "CSPM_RUNTIME_DIR": profile, "CSPM_DATA_DIR": profile,
                 "CSPM_LOG_DIR": Path(audit) / "runtime", "CSPM_EXPORT_DIR": profile / "exports",
                 "CSPM_MACHINE_ID_FILE": profile / "machine_id.json",
                 "CSPM_EXECUTABLE_ROOT": profile,
                 "PYTHONPYCACHEPREFIX": Path(audit) / "pycache",
                 "QTWEBENGINE_DICTIONARIES_PATH": profile / "dictionaries"}
    for name, path in overrides.items():
        checked_child(audit, path)
        if path.suffix != ".json":
            path.mkdir(parents=True, exist_ok=True)
        os.environ[name] = str(path)
    for name in ("OneDrive", "OneDriveConsumer", "OneDriveCommercial", "CSPM_PIXEL_DEPENDENCIES",
                 "CSPM_WRITE_BYTECODE"):
        os.environ.pop(name, None)
    sys.pycache_prefix = str(Path(audit) / "pycache")
    return profile, overrides


def fixture_source(source, root, audit, profile):
    """Remove the historical private-data bootstrap and automatic deadline."""
    start = source.index("ROOT = Path(__file__)")
    end = source.index('sys.path.insert(0, str(ROOT / "src/python"))', start)
    source = source[:start] + (f"ROOT = Path({str(root)!r})\nAUDIT = Path({str(audit)!r})\n"
              f"PROFILE = Path({str(profile)!r})\n") + source[end:]
    anchor = "        QTimer.singleShot(120000, self.timeout)"
    if source.count(anchor) != 1:
        raise ValueError("Manual fixture startup anchor is ambiguous")
    return source.replace(anchor, "        # Manual fixture has no unattended assessment deadline.")


def shell_source(source):
    """Patch only the disposable copy; keep ordinary command guards verbatim."""
    anchor = '        if (userMoveInProgress || userResizeInProgress || systemMoveInProgress) return false;\n        if (uiMaximized) {'
    if source.count(anchor) != 1:
        raise ValueError("Manual control interception anchor is ambiguous")
    source = source.replace(anchor, anchor.split('\n        if (uiMaximized)')[0]
                            + '\n        return manualNativeFixture.requestToggle();\n        if (uiMaximized) {', 1)
    # Also fail closed for drag/alternate entrypoints that could select production.
    anchor = '    function beginProfessionalWindowTransition(kind, sourceRect, targetRect,\n            targetScreenOverride) {'
    if source.count(anchor) != 1:
        raise ValueError("Professional transition entry anchor is ambiguous")
    source = source.replace(anchor, anchor + '\n        manualNativeFixture.rejectAlternateCommand(); return false;', 1)
    start = source.index('    title: mainWin.detachedMode')
    end = source.index('\n', source.index(': "CSPM - Main Menu"', start))
    source = source[:start] + '    title: "' + LABEL + '"' + source[end:]
    # Keep a visible label in the actual captured workspace, not only the taskbar.
    root = '    id: mainWin\n'
    if source.count(root) != 1:
        raise ValueError("Manual shell root is ambiguous")
    label = '''    Rectangle {
        parent: mainWin.contentItem
        anchors.horizontalCenter: parent.horizontalCenter
        y: 60; width: Math.min(parent.width - 100, 610); height: 28; z: 100000
        color: "#7b3500"; radius: 4
        Text { anchors.centerIn: parent; text: "CSPM NATIVE MOTION VISUAL TEST — DISPOSABLE";
            color: "white"; font.pixelSize: 13; font.bold: true }
    }
'''
    return source.replace(root, root + label, 1)


def install_io_guard(root, audit):
    """Process-local Python I/O guard, alongside Qt/CSPM path redirection."""
    write_root = Path(audit).resolve()
    readable = tuple(Path(p).resolve() for p in (root, audit, sys.base_prefix, os.environ["SystemRoot"]))

    def allowed(path, write):
        if isinstance(path, int) or path is None:
            return True
        if os.fsdecode(path).lower() == os.devnull.lower():
            return True  # Windows discard device, not an application-state file.
        value = Path(os.fsdecode(path)).resolve()
        if not write and value.is_relative_to(Path(root).resolve()) and not value.is_relative_to(write_root):
            relative = value.relative_to(Path(root).resolve())
            if relative.parts[0] in ("data", "backups", "logs") or value.name == "user_settings.json":
                return False
        roots = (write_root,) if write else readable
        return any(value == item or value.is_relative_to(item) for item in roots)

    def guard(event, args):
        if event == "open":
            path, mode, flags = args
            write = (isinstance(mode, str) and any(c in mode for c in "wax+")) or bool(flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC))
            if not allowed(path, write):
                frame = sys._getframe(1)
                stack = []
                for _ in range(12):
                    if frame is None:
                        break
                    stack.append({"file": frame.f_code.co_filename, "line": frame.f_lineno,
                                  "function": frame.f_code.co_name})
                    frame = frame.f_back
                with (write_root / "io_rejections.jsonl").open("a", encoding="utf-8") as output:
                    output.write(json.dumps({"event": event, "path": os.fsdecode(path),
                                             "write": write, "stack": stack}) + "\n")
                raise PermissionError("Manual fixture denied I/O outside disposable/source/runtime roots")
        elif event in ("os.remove", "os.rmdir", "os.mkdir", "os.chmod", "os.utime"):
            if not allowed(args[0], True):
                raise PermissionError("Manual fixture denied mutation outside disposable root")
        elif event in ("os.rename", "os.link", "os.symlink"):
            if not all(allowed(path, True) for path in args[:2]):
                raise PermissionError("Manual fixture denied file transfer outside disposable root")
        elif event in ("winreg.SetValue", "winreg.DeleteValue", "winreg.DeleteKey", "winreg.CreateKey"):
            raise PermissionError("Manual fixture forbids registry changes")
        elif event == "subprocess.Popen":
            if (args[3] is not None or Path(args[2] or "").resolve() != Path(root).resolve()
                    or not permitted_briefing_worker(root, audit, args[0], args[1])):
                raise PermissionError("Manual fixture permits only its isolated synthetic briefing worker")
        elif event in ("socket.connect", "socket.bind", "os.system", "os.startfile"):
            raise PermissionError("Manual fixture forbids external actions")
    sys.addaudithook(guard)
    return guard


def permitted_briefing_worker(root, audit, executable, argv):
    """Allow the existing crash-isolated bootstrap only with owned input/output."""
    if isinstance(argv, str) and sys.platform == "win32":
        import ctypes
        from ctypes import wintypes
        count = ctypes.c_int()
        shell = ctypes.WinDLL("shell32")
        shell.CommandLineToArgvW.argtypes = [wintypes.LPCWSTR, ctypes.POINTER(ctypes.c_int)]
        shell.CommandLineToArgvW.restype = ctypes.POINTER(wintypes.LPWSTR)
        pointers = shell.CommandLineToArgvW(argv, ctypes.byref(count))
        if not pointers:
            return False
        try:
            argv = [pointers[index] for index in range(count.value)]
        finally:
            kernel = ctypes.WinDLL("kernel32")
            kernel.LocalFree.argtypes = [ctypes.c_void_p]
            kernel.LocalFree.restype = ctypes.c_void_p
            kernel.LocalFree(pointers)
    if not isinstance(argv, (list, tuple)) or len(argv) != 5:
        return False
    if ((executable is not None and Path(executable).resolve() != Path(sys.executable).resolve())
            or Path(argv[0]).resolve() != Path(sys.executable).resolve()
            or Path(argv[1]).resolve() != (Path(root) / "src/python/main.py").resolve()
            or argv[2] != "--startup-briefing-worker"):
        return False
    try:
        request_path, result_path = (checked_child(audit, value) for value in argv[3:])
        payload = json.loads(request_path.read_text(encoding="utf-8"))
        data = checked_child(audit, payload["dataDir"])
        return (Path(payload["root"]).resolve() == Path(root).resolve()
                and request_path.parent == result_path.parent == data.parent
                and (data / "CSPM.xlsm").is_file())
    except (ValueError, OSError, KeyError, TypeError):
        return False


def attach_status(app, source_sha, dll_hash, audit):
    """Small separate status window; its content never feeds presentation."""
    from PySide6.QtCore import QObject, Slot, Qt
    from PySide6.QtWidgets import QWidget, QVBoxLayout, QLabel, QPushButton

    class ManualController(QObject):
        @Slot(result=bool)
        def requestToggle(self):
            if app.finished or app.manual_busy or not app.manual_ready:
                return False
            if app.window.property("appStyle") != "Professional":
                self.rejectAlternateCommand()
                return False
            app.manual_busy = True
            app.manual_requested = True
            self.update("direction", "restore" if app.window.property("uiMaximized") else "maximize")
            for name in ("source preparation", "target readiness", "input restoration"):
                self.update(name, "Pending")
            self.update("actually used", "Awaiting native preparation; production fallback disabled")
            app.step()
            return True

        @Slot()
        def rejectAlternateCommand(self):
            self.update("rejection", "Alternate transition refused; use maximize/restore control")

        def update(self, name, value):
            self.values[name] = str(value)
            self.text.setText("\n".join(f"{key}: {value}" for key, value in self.values.items()))
            (Path(audit) / "manual_status.json").write_text(json.dumps(self.values, indent=2), encoding="utf-8")

        def event(self, name, data):
            if name == "native-clock-start":
                self.update("actually used", "Native experimental GPU composition")
                self.update("source preparation", "Accepted by existing native guards")
            elif name == "target-gpu-composition-committed":
                self.update("target readiness", "Imported within unchanged native deadline")
            elif name == "input-restored":
                self.update("input restoration", "Enabled + QML interactive" if data.get("nativeEnabled") and data.get("qmlInteractive") else "FAILED")
            elif name == "qualification-failure":
                self.update("rejection", str(data.get("category")) + ": " + str(data.get("error")))
                if "deadline" in str(data.get("error", "")).lower() or "target" in str(data.get("error", "")).lower():
                    self.update("target readiness", "REJECTED by unchanged guard")
                if self.values["source preparation"] == "Pending":
                    self.update("source preparation", "REJECTED before native motion")
                    self.update("actually used", "No transition — safe native rejection; no production fallback")
                else:
                    self.update("actually used", "Native experimental attempted — safely rejected; no production fallback")
                if "slot" in str(data.get("error", "")).lower():
                    self.update("presentation slot", "REJECTED by native 100 ms gate")

    controller = ManualController(app)
    controller.values = {"requested": "Native experimental GPU composition", "actually used": "Native path armed; no transition requested",
        "source SHA": source_sha, "bridge SHA-256": dll_hash, "direction": "None (manual only)",
        "source preparation": "Not requested", "target readiness": "Not requested",
        "presentation slot": "Not traced; native 100 ms gate unchanged",
        "rejection": "None; production fallback disabled", "input restoration": "Not requested",
        "external observation": "DISABLED", "qualification": "Visual evaluation only — NOT QUALIFIED"}
    panel = QWidget()
    panel.setWindowTitle("Native Motion Status — Disposable")
    layout = QVBoxLayout(panel)
    heading = QLabel(LABEL)
    heading.setStyleSheet("color: #a64100; font-weight: bold")
    layout.addWidget(heading)
    controller.text = QLabel()
    controller.text.setTextInteractionFlags(Qt.TextSelectableByMouse)
    layout.addWidget(controller.text)
    close = QPushButton("Close disposable fixture")
    close.clicked.connect(lambda: app.quit() if app.finished else app.complete())
    layout.addWidget(close)
    controller.panel = panel
    controller.update("requested", controller.values["requested"])
    app.engine.rootContext().setContextProperty("manualNativeFixture", controller)
    app.manual_controller = controller
    panel.resize(640, 360)
    area = app.window.screen().availableGeometry()
    panel.move(area.right() - 650, area.bottom() - 370)
    panel.show()
    app.window.setTitle(LABEL)
    app.manual_ready = True
    app.manual_busy = False
    # No request or automated transition is scheduled here.
    return controller
