from __future__ import annotations

import json
import os
import shutil
import socket
import zipfile
from pathlib import Path

import pytest

from services.paths import AppPaths
from services.sync_service import SyncService


@pytest.fixture(autouse=True)
def isolated_sync_profile(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # Template tests use the real immutable source root for approval evidence;
    # each test still needs independent ancestry and recovery state outside it.
    monkeypatch.setenv("CSPM_RUNTIME_DIR", str(tmp_path / "profile/runtime"))
    monkeypatch.setenv("CSPM_DATA_DIR", str(tmp_path / "profile/persistent"))


def _write_package(folder: Path, label: str) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    # Distinct synthetic bytes exercise synchronization contracts; approved
    # empty template identity is tested separately with real workbooks.
    (folder / "CSPM.xlsm").write_bytes((f"CSPM-{label}|".encode("utf-8")) * 7000)
    (folder / "Dockets.xlsm").write_bytes((f"DOCKETS-{label}|".encode("utf-8")) * 1200)


def _write_ooxml_workbook(path: Path, modified_at: str, payload: str = "same-data") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    core = (
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        '<dcterms:modified xsi:type="dcterms:W3CDTF">'
        f"{modified_at}"
        "</dcterms:modified><cp:lastModifiedBy>Cory Schneider</cp:lastModifiedBy>"
        "</cp:coreProperties>"
    )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as workbook:
        workbook.writestr("docProps/core.xml", core)
        workbook.writestr("xl/workbook.xml", f"<workbook><payload>{payload}</payload></workbook>")
        workbook.writestr("xl/vbaProject.bin", b"governed-macro-payload")
        padding = zipfile.ZipInfo("xl/media/test-padding.bin")
        padding.compress_type = zipfile.ZIP_STORED
        workbook.writestr(padding, b"x" * 60_000)


def _service(tmp_path: Path, name: str, cloud: Path, local: Path) -> SyncService:
    return SyncService(
        AppPaths(
            root=tmp_path / f"runtime-{name}",
            override_data_dir=local,
            override_master_dir=cloud,
        ),
        machine_identity_file=tmp_path / f"runtime-{name}" / SyncService.MACHINE_ID_FILE_NAME,
    )


def test_checkout_blocks_second_machine_then_publishes_verified_release(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local_a = tmp_path / "local-a"
    local_b = tmp_path / "local-b"
    _write_package(cloud, "initial")

    first = _service(tmp_path, "a", cloud, local_a)
    checked_out = first.checkout_from_cloud()
    assert checked_out["ok"]
    assert (cloud / first.LEASE_FILE_NAME).is_file()
    assert (local_a / "CSPM.xlsm").read_bytes() == (cloud / "CSPM.xlsm").read_bytes()

    second = _service(tmp_path, "b", cloud, local_b)
    blocked = second.checkout_from_cloud()
    assert not blocked["ok"]
    assert blocked["status"] == "checkout-held"
    with pytest.raises(PermissionError):
        second.assert_write_lease()

    _write_package(local_a, "published")
    published = first.publish_and_release()
    assert published["ok"]
    assert published["status"] == "push"
    release = published["release"]
    release_dir = Path(release["releasePath"])
    assert (release_dir / "release.json").is_file()
    assert (release_dir / "CSPM.xlsm").read_bytes() == (local_a / "CSPM.xlsm").read_bytes()
    assert not (cloud / first.LEASE_FILE_NAME).exists()
    assert (cloud / "CSPM.xlsm").read_bytes() == (local_a / "CSPM.xlsm").read_bytes()

    checked_out_second = second.checkout_from_cloud()
    assert checked_out_second["ok"]
    assert (local_b / "CSPM.xlsm").read_bytes() == (cloud / "CSPM.xlsm").read_bytes()
    assert second.publish_and_release()["ok"]


def test_publish_refuses_cloud_change_after_checkout_and_preserves_both_copies(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "base")
    service = _service(tmp_path, "a", cloud, local)
    assert service.checkout_from_cloud()["ok"]

    _write_package(local, "local-change")
    _write_package(cloud, "cloud-change")  # Simulates an out-of-band cloud edit.
    cloud_before_publish = (cloud / "CSPM.xlsm").read_bytes()
    local_before_publish = (local / "CSPM.xlsm").read_bytes()

    result = service.publish_and_release()
    assert not result["ok"]
    assert result["status"] == "conflict"
    assert (cloud / "CSPM.xlsm").read_bytes() == cloud_before_publish
    assert (local / "CSPM.xlsm").read_bytes() == local_before_publish
    assert not (cloud / service.LEASE_FILE_NAME).exists()


def test_checkout_aligns_modified_timestamp_only_difference_to_exact_cloud_hash(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_ooxml_workbook(cloud / "CSPM.xlsm", "2026-09-07T22:15:43Z")
    _write_ooxml_workbook(cloud / "Dockets.xlsm", "2026-09-07T22:15:43Z")
    service = _service(tmp_path, "a", cloud, local)
    assert service.checkout_from_cloud()["ok"]
    assert service.publish_and_release()["ok"]

    _write_ooxml_workbook(local / "CSPM.xlsm", "2026-09-07T23:05:06Z")
    _write_ooxml_workbook(cloud / "CSPM.xlsm", "2026-09-07T23:05:07Z")
    restarted = _service(tmp_path, "a", cloud, local)

    result = restarted.checkout_from_cloud()

    assert result["ok"]
    assert result["status"] == "metadata-aligned"
    assert result["metadataAligned"] == ["CSPM.xlsm"]
    assert (local / "CSPM.xlsm").read_bytes() == (cloud / "CSPM.xlsm").read_bytes()
    assert restarted._get_state()["files"]["CSPM.xlsm"]["baseHash"] == restarted._sha256(
        cloud / "CSPM.xlsm"
    )
    assert restarted.publish_and_release()["ok"]


def test_modified_timestamp_tolerance_never_hides_workbook_payload_change(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_ooxml_workbook(cloud / "CSPM.xlsm", "2026-09-07T22:15:43Z")
    _write_ooxml_workbook(cloud / "Dockets.xlsm", "2026-09-07T22:15:43Z")
    service = _service(tmp_path, "a", cloud, local)
    assert service.checkout_from_cloud()["ok"]
    assert service.publish_and_release()["ok"]

    _write_ooxml_workbook(local / "CSPM.xlsm", "2026-09-07T23:05:06Z", payload="local-data")
    _write_ooxml_workbook(cloud / "CSPM.xlsm", "2026-09-07T23:05:07Z", payload="cloud-data")
    restarted = _service(tmp_path, "a", cloud, local)

    result = restarted.checkout_from_cloud()

    assert not result["ok"]
    assert result["status"] == "conflict"
    assert (local / "CSPM.xlsm").read_bytes() != (cloud / "CSPM.xlsm").read_bytes()
    assert not (cloud / restarted.LEASE_FILE_NAME).exists()


def test_unknown_existing_difference_is_a_conflict_not_an_automatic_copy(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "cloud")
    _write_package(local, "local")
    service = _service(tmp_path, "a", cloud, local)

    result = service.checkout_from_cloud()
    assert not result["ok"]
    assert result["status"] == "conflict"
    assert (cloud / "CSPM.xlsm").read_bytes() != (local / "CSPM.xlsm").read_bytes()
    assert not (cloud / service.LEASE_FILE_NAME).exists()


def test_fresh_blank_seed_package_pulls_the_complete_cloud_pair(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "cloud")
    local.mkdir()
    # Both existing files must be approved empty templates. An arbitrary or
    # populated workbook must never be discarded as part of a fresh replica.
    project_root = Path(__file__).resolve().parents[1]
    shutil.copyfile(project_root / "src/templates/CSPM.xlsm", local / "CSPM.xlsm")
    shutil.copyfile(project_root / "src/templates/Dockets.xlsm", local / "Dockets.xlsm")
    service = SyncService(AppPaths(project_root, local, cloud), machine_identity_file=tmp_path / "machine.json")

    result = service.checkout_from_cloud()
    assert result["ok"]
    assert result["status"] == "pull"
    assert (local / "CSPM.xlsm").read_bytes() == (cloud / "CSPM.xlsm").read_bytes()
    assert (local / "Dockets.xlsm").read_bytes() == (cloud / "Dockets.xlsm").read_bytes()
    assert service.publish_and_release()["ok"]


@pytest.mark.parametrize("cspm_present", [True, False])
def test_approved_or_missing_cspm_does_not_discard_populated_dockets(tmp_path: Path, cspm_present: bool) -> None:
    from openpyxl import load_workbook

    project_root = Path(__file__).resolve().parents[1]
    cloud, local = tmp_path / "cloud", tmp_path / "local"
    _write_package(cloud, "governed")
    local.mkdir()
    if cspm_present:
        shutil.copyfile(project_root / "src/templates/CSPM.xlsm", local / "CSPM.xlsm")
    dockets = local / "Dockets.xlsm"
    shutil.copyfile(project_root / "src/templates/Dockets.xlsm", dockets)
    workbook = load_workbook(dockets, keep_vba=True)
    try:
        workbook["Dockets"].cell(2, 1).value = "SYNTHETIC LOCAL DOCKET MUST BE PRESERVED"
        workbook.save(dockets)
    finally:
        for attribute in ("vba_archive", "_archive"):
            archive = getattr(workbook, attribute, None)
            if archive is not None:
                archive.close()
        workbook.close()
    before = {str(path): path.read_bytes() for directory in (cloud, local)
              for path in directory.glob("*.xlsm")}
    service = SyncService(AppPaths(project_root, local, cloud), machine_identity_file=tmp_path / "machine.json")
    assert not service._is_blank_seed("Dockets.xlsm", dockets)
    result = service.checkout_from_cloud()
    assert not result["ok"] and result["status"] == "conflict"
    assert "Dockets.xlsm" in result["files"]
    assert {str(path): path.read_bytes() for directory in (cloud, local)
            for path in directory.glob("*.xlsm")} == before
    assert not (cloud / service.LEASE_FILE_NAME).exists()


def test_small_populated_workbook_saves_publishes_and_survives_recheckout(tmp_path: Path) -> None:
    from repositories.excel_repo import ExcelRepo
    from domain import schema_constants as sc

    project_root = Path(__file__).resolve().parents[1]
    cloud, local = tmp_path / "cloud", tmp_path / "local"
    cloud.mkdir()
    for name in SyncService._FILES:
        shutil.copyfile(project_root / "src/templates" / name, cloud / name)

    service = SyncService(AppPaths(project_root, local, cloud), machine_identity_file=tmp_path / "machine.json")
    assert service.checkout_from_cloud()["ok"]
    repo = ExcelRepo(service.paths, write_guard=service.assert_write_lease)
    result = repo.add_time_entry({"date": "2026-10-12", "clientName": "Synthetic Client",
                                 "matterName": "Synthetic Matter", "description": "Synthetic small practice docket",
                                 "hours": 1, "rate": 10, "sharePct": 100})
    assert result["ok"]
    assert (local / "CSPM.xlsm").stat().st_size < 75_000
    assert not service._is_blank_seed("CSPM.xlsm", local / "CSPM.xlsm")

    published = service.publish_and_release()
    assert published["ok"] and published["status"] == "push"
    release = Path(published["release"]["releasePath"])
    manifest = json.loads((release / service.RELEASE_MANIFEST_NAME).read_text())
    for name in SyncService._FILES:
        assert service._sha256(release / name) == service._sha256(cloud / name) == manifest["files"][name]
    assert not (cloud / service.LEASE_FILE_NAME).exists()

    restarted = SyncService(service.paths, machine_identity_file=tmp_path / "machine.json")
    assert restarted.checkout_from_cloud()["ok"]
    reopened = ExcelRepo(restarted.paths, write_guard=restarted.assert_write_lease)
    assert any(row.get(sc.COL_TIME_ENTRY_ID) == result["entryId"] for row in reopened._read_table_rows(sc.TBL_TIME))
    assert restarted.publish_and_release()["ok"]


def test_approved_empty_first_checkout_closes_without_false_conflict(tmp_path: Path) -> None:
    project_root = Path(__file__).resolve().parents[1]
    cloud, local = tmp_path / "cloud", tmp_path / "local"
    cloud.mkdir()
    for name in SyncService._FILES:
        shutil.copyfile(project_root / "src/templates" / name, cloud / name)
    service = SyncService(AppPaths(project_root, local, cloud), machine_identity_file=tmp_path / "machine.json")
    assert service.checkout_from_cloud()["ok"]
    assert service._is_blank_seed("CSPM.xlsm", local / "CSPM.xlsm")
    result = service.publish_and_release()
    assert result["ok"] and result["status"] == "unchanged"
    assert not (cloud / service.LEASE_FILE_NAME).exists()


def test_approved_seed_does_not_override_existing_common_ancestor(tmp_path: Path) -> None:
    project_root = Path(__file__).resolve().parents[1]
    cloud, local = tmp_path / "cloud", tmp_path / "local"
    _write_package(cloud, "initial")
    service = SyncService(AppPaths(project_root, local, cloud), machine_identity_file=tmp_path / "machine.json")
    assert service.checkout_from_cloud()["ok"]
    assert service.publish_and_release()["ok"]
    # Simulate divergent local reset and independent shared edit. The local
    # exact empty template does not authorize discarding known offline changes.
    for name in SyncService._FILES:
        shutil.copyfile(project_root / "src/templates" / name, local / name)
    _write_package(cloud, "different-governed-revision")
    before = {str(path): path.read_bytes() for directory in (cloud, local)
              for path in directory.glob("*.xlsm")}
    restarted = SyncService(service.paths, machine_identity_file=tmp_path / "machine.json")
    assert restarted._is_blank_seed("CSPM.xlsm", local / "CSPM.xlsm")
    result = restarted.checkout_from_cloud()
    assert not result["ok"] and result["status"] == "conflict"
    assert {str(path): path.read_bytes() for directory in (cloud, local)
            for path in directory.glob("*.xlsm")} == before
    assert not (cloud / service.LEASE_FILE_NAME).exists()


def test_small_unknown_workbook_is_preserved_as_conflict(tmp_path: Path) -> None:
    cloud, local = tmp_path / "cloud", tmp_path / "local"
    _write_package(cloud, "governed")
    local.mkdir()
    (local / "CSPM.xlsm").write_bytes(b"small unknown populated practice")
    (local / "Dockets.xlsm").write_bytes(b"small unknown dockets")
    before = {name: (local / name).read_bytes() for name in SyncService._FILES}
    service = _service(tmp_path, "a", cloud, local)
    result = service.checkout_from_cloud()
    assert not result["ok"] and result["status"] == "conflict"
    assert {name: (local / name).read_bytes() for name in SyncService._FILES} == before
    assert not (cloud / service.LEASE_FILE_NAME).exists()


def test_approved_seed_allows_only_modified_metadata_difference(tmp_path: Path) -> None:
    project_root = Path(__file__).resolve().parents[1]
    seed = tmp_path / "CSPM.xlsm"
    template = project_root / "src/templates/CSPM.xlsm"
    with zipfile.ZipFile(template) as source, zipfile.ZipFile(seed, "w") as target:
        for item in source.infolist():
            payload = source.read(item)
            if item.filename == "docProps/core.xml":
                # Change specifically the modified timestamp, preserving creator/created.
                from xml.etree import ElementTree
                properties = ElementTree.fromstring(source.read(item))
                modified = properties.find("{http://purl.org/dc/terms/}modified")
                assert modified is not None
                modified.text = "2026-10-12T00:00:00Z"
                payload = ElementTree.tostring(properties, encoding="utf-8")
            target.writestr(item, payload)
    service = SyncService(AppPaths(project_root, tmp_path), machine_identity_file=tmp_path / "machine.json")
    assert service._sha256(seed) != service._sha256(template)
    assert service._is_blank_seed("CSPM.xlsm", seed)
    with zipfile.ZipFile(seed) as source:
        payloads = [(item, source.read(item)) for item in source.infolist()]
    with zipfile.ZipFile(seed, "w") as target:
        for item, payload in payloads:
            if item.filename == "docProps/core.xml":
                payload = payload.replace(b"System", b"Changed author")
            target.writestr(item, payload)
    assert not service._is_blank_seed("CSPM.xlsm", seed)


@pytest.mark.parametrize("failure", ["missing-approval", "malformed-approval", "unapproved", "wrong-template-hash"])
def test_unverified_template_never_authorizes_seed_replacement(tmp_path: Path, failure: str) -> None:
    project_root = Path(__file__).resolve().parents[1]
    template_dir = tmp_path / "runtime/src/templates"
    template_dir.mkdir(parents=True)
    template = template_dir / "CSPM.xlsm"
    shutil.copyfile(project_root / "src/templates/CSPM.xlsm", template)
    seed = tmp_path / "CSPM.xlsm"
    shutil.copyfile(template, seed)
    approval_file = template_dir / "template_approval.json"
    approval = json.loads((project_root / "src/templates/template_approval.json").read_text())
    if failure == "malformed-approval":
        approval_file.write_text("not json")
    elif failure != "missing-approval":
        if failure == "unapproved":
            approval["templates"][0]["human_approval_status"] = "PENDING"
        elif failure == "wrong-template-hash":
            approval["templates"][0]["sha256"] = "0" * 64
        approval_file.write_text(json.dumps(approval))
    service = SyncService(AppPaths(tmp_path / "runtime", tmp_path), machine_identity_file=tmp_path / "machine.json")
    assert not service._is_blank_seed("CSPM.xlsm", seed)


def test_existing_cloud_setup_never_overwrites_a_different_local_package(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "cloud")
    _write_package(local, "local")
    service = _service(tmp_path, "a", cloud, local)

    result = service.initialize_shared_source(local)
    assert not result["ok"]
    assert result["status"] == "conflict"
    assert (cloud / "CSPM.xlsm").read_bytes() != (local / "CSPM.xlsm").read_bytes()


def test_repository_write_guard_is_available_only_during_checkout(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "base")
    service = _service(tmp_path, "a", cloud, local)

    with pytest.raises(PermissionError):
        service.assert_write_lease()
    assert service.checkout_from_cloud()["ok"]
    service.assert_write_lease()
    assert service.publish_and_release()["ok"]
    with pytest.raises(PermissionError):
        service.assert_write_lease()


def test_same_installation_dead_process_checkout_is_recovered_automatically(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "base")
    first = _service(tmp_path, "a", cloud, local)
    stale_lease = {
        "schemaVersion": 1,
        "machineId": first._machine_identity(),
        "checkoutId": "dead-checkout",
        "processId": 424242,
        "computerName": socket.gethostname(),
        "checkedOutAtUtc": "2026-08-14T00:00:00+00:00",
        "purpose": "CSPM exclusive workbook write checkout",
    }
    (cloud / first.LEASE_FILE_NAME).write_text(json.dumps(stale_lease), encoding="utf-8")

    restarted = _service(tmp_path, "a", cloud, local)
    monkeypatch.setattr(restarted, "_local_process_is_running", lambda _pid: False)

    acquired = restarted._acquire_checkout_lease()

    assert acquired["ok"]
    assert acquired["status"] == "checked-out"
    recovery = acquired["recovery"]
    assert recovery["status"] == "recovered-abandoned-checkout"
    audit_path = Path(recovery["recoveryAuditPath"])
    assert audit_path.is_file()
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    assert audit["abandonedLease"] == stale_lease
    assert json.loads((cloud / first.LEASE_FILE_NAME).read_text(encoding="utf-8"))["checkoutId"] == restarted._checkout_id
    assert restarted._release_checkout_lease()["ok"]


def test_auto_recovery_never_removes_another_installations_dead_process_checkout(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "base")
    owner = _service(tmp_path, "owner", cloud, local)
    other = _service(tmp_path, "other", cloud, tmp_path / "other-local")
    foreign_lease = {
        "schemaVersion": 1,
        "machineId": owner._machine_identity(),
        "checkoutId": "foreign-dead-checkout",
        "processId": 424242,
        "computerName": socket.gethostname(),
        "checkedOutAtUtc": "2026-08-14T00:00:00+00:00",
        "purpose": "CSPM exclusive workbook write checkout",
    }
    (cloud / owner.LEASE_FILE_NAME).write_text(json.dumps(foreign_lease), encoding="utf-8")
    monkeypatch.setattr(other, "_local_process_is_running", lambda _pid: False)

    blocked = other._acquire_checkout_lease()

    assert not blocked["ok"]
    assert blocked["status"] == "checkout-held"
    assert json.loads((cloud / owner.LEASE_FILE_NAME).read_text(encoding="utf-8")) == foreign_lease


def test_source_and_packaged_roots_share_stable_workstation_identity(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "base")
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "local-app-data"))

    source = SyncService(
        AppPaths(root=tmp_path / "source", override_data_dir=local, override_master_dir=cloud)
    )
    packaged = SyncService(
        AppPaths(root=tmp_path / "package", override_data_dir=local, override_master_dir=cloud)
    )
    assert source._machine_identity() == packaged._machine_identity()

    stale_lease = {
        "schemaVersion": 1,
        "machineId": source._machine_identity(),
        "checkoutId": "dead-source-checkout",
        "processId": 424242,
        "computerName": socket.gethostname(),
        "checkedOutAtUtc": "2026-09-08T00:00:00+00:00",
        "purpose": "CSPM exclusive workbook write checkout",
    }
    (cloud / source.LEASE_FILE_NAME).write_text(json.dumps(stale_lease), encoding="utf-8")
    monkeypatch.setattr(packaged, "_local_process_is_running", lambda _pid: False)

    acquired = packaged._acquire_checkout_lease()

    assert acquired["ok"]
    assert acquired["recovery"]["status"] == "recovered-abandoned-checkout"
    assert packaged._release_checkout_lease()["ok"]


def test_auto_recovery_does_not_reclaim_a_same_installation_live_process_checkout(tmp_path: Path) -> None:
    cloud = tmp_path / "cloud"
    local = tmp_path / "local"
    _write_package(cloud, "base")
    owner = _service(tmp_path, "a", cloud, local)
    live_lease = {
        "schemaVersion": 1,
        "machineId": owner._machine_identity(),
        "checkoutId": "live-checkout",
        "processId": os.getpid(),
        "computerName": socket.gethostname(),
        "checkedOutAtUtc": "2026-08-14T00:00:00+00:00",
        "purpose": "CSPM exclusive workbook write checkout",
    }
    (cloud / owner.LEASE_FILE_NAME).write_text(json.dumps(live_lease), encoding="utf-8")
    restarted = _service(tmp_path, "a", cloud, local)

    blocked = restarted._acquire_checkout_lease()

    assert not blocked["ok"]
    assert blocked["status"] == "checkout-held"
    assert json.loads((cloud / owner.LEASE_FILE_NAME).read_text(encoding="utf-8")) == live_lease
