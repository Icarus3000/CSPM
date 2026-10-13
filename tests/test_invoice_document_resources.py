"""Invoice resources stay in the package when the writable profile moves."""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document
from PySide6.QtCore import QCoreApplication


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src" / "python"))

from services.invoice_document_service import InvoiceDocumentService
from services import invoice_document_service
from services.paths import AppPaths


PAYLOAD = {
    "invoice_number_raw": "TEST-2026-001",
    "billing_client_name": "Synthetic Invoice Client",
    "billing_client_address_lines": ["1 Example Street", "Toronto, ON"],
    "date": "October 13, 2026",
    "is_draft": True,
    "is_flat_fee": False,
    "hourly_matters": [],
    "gross_professional_fees": 100.0,
    "total_fees": 100.0,
    "net_professional_fees": 100.0,
    "discount_amount": 0.0,
    "agency_split_amount": 0.0,
    "agency_split_percent": 0.0,
    "total_tax": 13.0,
    "total_due": 113.0,
    "disbursement_total": 0.0,
    "total_hours": 1.0,
    "show_total_hours": True,
    "has_any_flat_fee": False,
}


def _separate_runtime(tmp_path, monkeypatch):
    runtime = tmp_path / "runtime"
    runtime.mkdir()
    for variable in (
        "CSPM_EXECUTABLE_ROOT", "CSPM_RUNTIME_DIR", "CSPM_DATA_DIR", "LOCALAPPDATA"
    ):
        monkeypatch.setenv(variable, str(runtime))
    # An old profile may contain a copied or stale layout. It must never win
    # over the immutable layout supplied with this application version.
    stale_templates = runtime / "src" / "templates" / "invoices"
    stale_templates.mkdir(parents=True)
    (stale_templates / "Concept_A2.html").write_text("STALE PROFILE INVOICE")
    monkeypatch.chdir(runtime)
    return runtime


@pytest.mark.parametrize("frozen", [False, True])
def test_default_html_layout_ignores_profile_and_working_directory(
    tmp_path, monkeypatch, frozen
):
    _separate_runtime(tmp_path, monkeypatch)
    resource_root = PROJECT_ROOT
    if frozen:
        resource_root = tmp_path / "package" / "_internal"
        templates = resource_root / "src" / "templates" / "invoices"
        templates.mkdir(parents=True)
        shutil.copyfile(
            PROJECT_ROOT / "src" / "templates" / "invoices" / "Concept_A2.html",
            templates / "Concept_A2.html",
        )
        monkeypatch.setattr(sys, "_MEIPASS", str(resource_root), raising=False)
    monkeypatch.setattr(sys, "frozen", frozen, raising=False)

    service = InvoiceDocumentService()
    html = service.generate_html("Concept_A2", PAYLOAD)

    assert Path(service.templates_dir) == resource_root / "src" / "templates" / "invoices"
    assert "Synthetic Invoice Client" in html
    assert "TEST-2026-001" in html
    assert "$113.00" in html
    assert "STALE PROFILE INVOICE" not in html


def test_frozen_docx_layout_and_export_use_separate_locations(tmp_path, monkeypatch):
    runtime = _separate_runtime(tmp_path, monkeypatch)
    resource_root = tmp_path / "package" / "_internal"
    templates = resource_root / "src" / "templates" / "invoices"
    templates.mkdir(parents=True)
    template = Document()
    template.add_paragraph("Invoice {{ invoice_number_raw }} for {{ billing_client_name }}")
    source_path = templates / "Fixture.docx"
    template.save(source_path)
    original = source_path.read_bytes()
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "_MEIPASS", str(resource_root), raising=False)
    output = runtime / "invoice.docx"

    result = InvoiceDocumentService().generate_docx("Fixture", PAYLOAD, str(output))

    assert result == str(output)
    assert Document(output).paragraphs[0].text == (
        "Invoice TEST-2026-001 for Synthetic Invoice Client"
    )
    assert source_path.read_bytes() == original


def test_explicit_template_directory_remains_supported(tmp_path):
    (tmp_path / "Custom.html").write_text("Invoice {{ invoice_number_raw }}")
    service = InvoiceDocumentService(str(tmp_path))

    assert service.generate_html("Custom", PAYLOAD) == "Invoice TEST-2026-001"


@pytest.mark.parametrize("layout", ["Concept_A2", "Concept_B", "Concept_C", "Concept_A_Agency"])
def test_unavailable_word_layout_recommends_pdf_before_opening_output(
    tmp_path, monkeypatch, layout
):
    output = tmp_path / "existing.docx"
    output.write_bytes(b"PRESERVE EXISTING EXPORT")
    loader = Mock(side_effect=AssertionError("Missing layout must not enter docxtpl"))
    monkeypatch.setattr(invoice_document_service, "DocxTemplate", loader)

    with pytest.raises(ValueError, match="Use Export PDF"):
        InvoiceDocumentService().generate_docx(layout, PAYLOAD, str(output))

    loader.assert_not_called()
    assert output.read_bytes() == b"PRESERVE EXISTING EXPORT"


def test_application_constructs_invoice_service_from_resources(tmp_path, monkeypatch):
    from backend import native_candidate
    from backend.app_controller import AppController

    application = QCoreApplication.instance() or QCoreApplication([])
    runtime = _separate_runtime(tmp_path, monkeypatch)
    monkeypatch.setattr(sys, "frozen", False, raising=False)
    monkeypatch.setattr(native_candidate, "_profile", None)
    monkeypatch.setenv("CSPM_AUTOBACKUP_MINUTES", "0")
    data = runtime / "practice"
    data.mkdir()
    # Constructor-only fixture: no real workbook is opened or changed.
    for name in ("CSPM.xlsm", "Dockets.xlsm"):
        (data / name).write_bytes(b"")
    (runtime / "user_settings.json").write_text(json.dumps({
        "localDataDir": str(data),
        "masterDataDir": "",
        "autoBackupMinutes": 0,
        "keepTrayAlive": False,
    }))
    paths = AppPaths(runtime, override_data_dir=data)
    monkeypatch.setattr(paths, "workbook_path", lambda: runtime / "not-opened.xlsm")
    controller = AppController(paths=paths, defer_settings_load=True)
    try:
        html = controller._invoice_doc_svc.generate_html("Concept_A2", PAYLOAD)
        assert "Synthetic Invoice Client" in html
        assert "$113.00" in html
        assert "STALE PROFILE INVOICE" not in html
    finally:
        controller._auto_backup_timer.stop()
        controller._startup_briefing_process_poll_timer.stop()
        application.processEvents()
