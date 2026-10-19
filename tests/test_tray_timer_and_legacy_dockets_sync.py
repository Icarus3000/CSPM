from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from PySide6.QtCore import QCoreApplication

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src/python"))

from backend.controllers.tray_controller import TrayController
from domain import schema_constants as sc
from repositories.excel_repo import ExcelRepo, TBL_CLIENTS, TBL_MATTERS, TBL_PARENTS
from services.legacy_dockets_sync_service import (
    LEGACY_DOCKETS_COLUMNS,
    LEGACY_DOCKETS_SHEET,
    LEGACY_DOCKETS_TABLE,
    LegacyDocketsSyncService,
)
from services.paths import AppPaths


def _create_mock_legacy_dockets_workbook(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = LEGACY_DOCKETS_SHEET
    ws.append(list(LEGACY_DOCKETS_COLUMNS))
    ws.append(
        [
            dt.datetime(2026, 10, 15, 0, 0),
            "Existing Seed Client",
            None,
            "Seed docket row",
            1.0,
            475,
            1,
            "=([@[Time (in hrs)]]*[@[Hourly Rate/Flat Fee]])*[@Percentage]",
            "=[@[Amount to CS]]*1.13",
            None,
            "SEED-MAT-26-0001",
            3600,
            "Hourly",
        ]
    )
    ws.cell(row=2, column=1).number_format = "[$-409]mmmm\\ d\\,\\ yyyy;@"
    ws.cell(row=2, column=5).number_format = "0.00"
    ws.cell(row=2, column=6).number_format = "$#,##0.00"
    ws.cell(row=2, column=7).number_format = "0%"
    ws.cell(row=2, column=8).number_format = "_($* #,##0.00_);_($* (#,##0.00);_($* \"-\"??_);_(@_)"
    ws.cell(row=2, column=9).number_format = "$#,##0.00"

    tbl = Table(displayName=LEGACY_DOCKETS_TABLE, ref="A1:M2")
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tbl)
    wb.save(path)
    wb.close()
    return path


def _create_isolated_cspm_repo(tmp_path: Path) -> tuple[AppPaths, ExcelRepo, Path]:
    paths = AppPaths(tmp_path, override_data_dir=tmp_path / "data")
    legacy_path = _create_mock_legacy_dockets_workbook(paths.dockets_workbook_path())

    repo = ExcelRepo(paths)
    repo.ensure_schema()

    repo._upsert_row_by_key(
        TBL_PARENTS,
        sc.COL_PARENT_ID,
        "PAR-001",
        {
            sc.COL_PARENT_ID: "PAR-001",
            sc.COL_PARENT_NAME: "Megatech Group",
            sc.COL_PARENT_DEF_RATE: 475.0,
            sc.COL_PARENT_DEF_SHARE: 100.0,
            sc.COL_PARENT_ACTIVE: True,
        },
    )
    repo._upsert_row_by_key(
        TBL_CLIENTS,
        sc.COL_CLIENT_ID,
        "CLI-001",
        {
            sc.COL_CLIENT_ID: "CLI-001",
            sc.COL_CLIENT_NAME: "Megatech Integrated Services Ltd.",
            sc.COL_CLIENT_ACTIVE: True,
        },
    )
    repo._upsert_row_by_key(
        TBL_MATTERS,
        sc.COL_MATTER_ID,
        "MAT-001",
        {
            sc.COL_MATTER_ID: "MAT-001",
            sc.COL_MATTER_NUMBER: "MEGA-COR-26-0093",
            sc.COL_MATTER_CLIENT_ID: "CLI-001",
            sc.COL_MATTER_CLIENT_NAME: "Megatech Integrated Services Ltd.",
            sc.COL_MATTER_PARENT_ID: "PAR-001",
            sc.COL_MATTER_PARENT_NAME: "Megatech Group",
            sc.COL_MATTER_NAME: "Corporate Reorganization",
            sc.COL_MATTER_STATUS: "Open",
            sc.COL_MATTER_DEF_RATE: 475.0,
            sc.COL_MATTER_DEF_SHARE: 100.0,
        },
    )
    return paths, repo, legacy_path


def test_main_timer_syncs_to_tray_indicator_and_minimalist_popup(tmp_path: Path) -> None:
    _app = QCoreApplication.instance() or QCoreApplication([])
    mock_icon = MagicMock()
    mock_app_controller = SimpleNamespace(_excel_repo=None, _theme_name="Dark", _tray_icon=mock_icon)
    tray = TrayController(mock_app_controller)

    # Initially idle (no running timer)
    assert tray.hasRunningTimer is False
    assert tray.hasActiveOrPausedTimer is False

    # 1. Start main timer from main app
    tray.syncMainTimerState(
        {
            "ownerId": "TDV_TEST_1",
            "isRunning": True,
            "elapsedSeconds": 125,
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093 - Corporate Reorganization",
            "description": "Drafting s. 85 rollover resolution",
        }
    )

    assert tray.hasRunningTimer is True
    assert tray.hasActiveOrPausedTimer is True
    assert tray.hasDisplayEntry is True
    assert tray.displayStatus == "RUNNING"
    assert tray.displayElapsedSeconds == 125
    assert tray.displayElapsedFormatted == "00:02:05"
    assert tray.displayClientName == "Megatech Integrated Services Ltd."
    assert tray.displayMatterName == "MEGA-COR-26-0093 — Corporate Reorganization"
    assert tray.displayDescription == "Drafting s. 85 rollover resolution"
    assert "RUNNING (00:02:05)" in tray.trayTooltipText
    assert "Megatech Integrated Services Ltd." in tray.trayTooltipText
    mock_icon.setToolTip.assert_called_with(tray.trayTooltipText)

    # 2. Simulate 1-second tick update
    tray._timers[0]._state.started_at_monotonic -= 1.0
    tray.update_timers()
    assert tray.displayElapsedSeconds >= 126
    assert tray.displayElapsedFormatted == "00:02:06"
    assert "00:02:06" in tray.trayTooltipText

    # 3. Edit description from tray popup and verify signal emission to main app
    received_commands: list[dict] = []
    tray.mainTimerCommandRequested.connect(lambda cmd: received_commands.append(dict(cmd)))

    tray.updatePrimaryDescription("Updated narrative from system tray mini-popup")
    assert tray.displayDescription == "Updated narrative from system tray mini-popup"
    assert received_commands[-1]["action"] == "description"
    assert received_commands[-1]["description"] == "Updated narrative from system tray mini-popup"

    # 4. Pause from main app
    tray.syncMainTimerState(
        {
            "ownerId": "TDV_TEST_1",
            "isRunning": False,
            "elapsedSeconds": 180,
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093 - Corporate Reorganization",
            "description": "Updated narrative from system tray mini-popup",
        }
    )
    assert tray.hasRunningTimer is False
    assert tray.hasActiveOrPausedTimer is True
    assert tray.displayStatus == "PAUSED"
    assert tray.displayElapsedFormatted == "00:03:00"
    assert "PAUSED (00:03:00)" in tray.trayTooltipText

    # 5. Stop and save docket -> popup transitions to LATEST entry when no timer is running
    tray.onDocketSaved(
        {
            "entryId": "TM-0001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterName": "MEGA-COR-26-0093 - Corporate Reorganization",
            "description": "Final saved docket description",
            "rawSeconds": 1800,
            "hours": 0.5,
            "date": "2026-10-18",
            "stoppedTimer": True,
        }
    )
    assert tray.hasRunningTimer is False
    assert tray.hasActiveOrPausedTimer is False
    assert tray.hasDisplayEntry is True
    assert tray.displayStatus == "LATEST"
    assert tray.displayElapsedFormatted == "00:30:00"
    assert tray.displayClientName == "Megatech Integrated Services Ltd."
    assert tray.displayMatterName == "MEGA-COR-26-0093 — Corporate Reorganization"
    assert tray.displayDescription == "Final saved docket description"

    # 6. Verify compact minimalist flyout dimensions (340x218)
    assert TrayController.FLYOUT_WIDTH == 340
    assert TrayController.FLYOUT_HEIGHT == 218

    tray.shutdown()


def test_saved_and_updated_dockets_auto_sync_to_legacy_dockets_xlsm(tmp_path: Path) -> None:
    paths, repo, legacy_path = _create_isolated_cspm_repo(tmp_path)

    # 1. Add a new time entry in CSPM
    add_res = repo.add_time_entry(
        {
            "date": "2026-10-18",
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093",
            "parentId": "PAR-001",
            "description": "Prepare share subscription and set-off resolutions",
            "rawSeconds": 5400,
            "hours": 1.5,
            "useExactHours": True,
            "clientRate": 475.0,
            "sharePct": 100.0,
            "status": "Draft",
        }
    )
    entry_id = add_res["entryId"]
    assert entry_id

    # Verify Table1_1 in mock Dockets.xlsm has the new row and updated ref
    wb = openpyxl.load_workbook(legacy_path, keep_vba=True, data_only=False)
    ws = wb[LEGACY_DOCKETS_SHEET]
    assert ws.tables[LEGACY_DOCKETS_TABLE].ref == "A1:M3"
    assert ws.cell(row=3, column=1).value == dt.datetime(2026, 10, 18, 0, 0)
    assert ws.cell(row=3, column=1).number_format == "[$-409]mmmm\\ d\\,\\ yyyy;@"
    assert ws.cell(row=3, column=2).value == "Megatech Group"
    assert ws.cell(row=3, column=3).value == "Megatech Integrated Services Ltd."
    assert ws.cell(row=3, column=4).value == "Prepare share subscription and set-off resolutions"
    assert float(ws.cell(row=3, column=5).value) == 1.5
    assert float(ws.cell(row=3, column=6).value) == 475.0
    assert float(ws.cell(row=3, column=7).value) == 1.0
    assert ws.cell(row=3, column=7).number_format == "0%"
    assert str(ws.cell(row=3, column=8).value).startswith("=")
    assert str(ws.cell(row=3, column=9).value).startswith("=")
    assert ws.cell(row=3, column=10).value is None
    assert ws.cell(row=3, column=11).value == "MEGA-COR-26-0093"
    assert int(ws.cell(row=3, column=12).value) == 5400
    assert ws.cell(row=3, column=13).value == "Hourly"
    LegacyDocketsSyncService._close_workbook(wb)

    # 2. Update the existing time entry in CSPM and verify Table1_1 updates in place without duplicating
    updated = repo.update_time_entry(
        entry_id,
        {
            "date": "2026-10-18",
            "hours": 2.2,
            "rawSeconds": 7920,
            "description": "Prepare and finalize share subscription and set-off resolutions",
            "invoiceRef": "26-0042",
        },
    )
    assert updated.get("ok") is True

    wb2 = openpyxl.load_workbook(legacy_path, keep_vba=True, data_only=False)
    ws2 = wb2[LEGACY_DOCKETS_SHEET]
    assert ws2.tables[LEGACY_DOCKETS_TABLE].ref == "A1:M3"
    assert ws2.max_row == 3
    assert ws2.cell(row=3, column=4).value == "Prepare and finalize share subscription and set-off resolutions"
    assert float(ws2.cell(row=3, column=5).value) == 2.2
    assert ws2.cell(row=3, column=10).value == "26-0042"
    assert ws2.cell(row=3, column=11).value == "MEGA-COR-26-0093"
    assert int(ws2.cell(row=3, column=12).value) == 7920
    LegacyDocketsSyncService._close_workbook(wb2)


def test_legacy_dockets_lock_resilience_queues_and_flushes(tmp_path: Path) -> None:
    paths, repo, legacy_path = _create_isolated_cspm_repo(tmp_path)

    # Simulate Excel holding a lock file (~$Dockets.xlsm)
    lock_file = legacy_path.parent / f"~${legacy_path.name}"
    lock_file.write_bytes(b"locked-by-excel")

    add_res = repo.add_time_entry(
        {
            "date": "2026-10-18",
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093",
            "parentId": "PAR-001",
            "description": "Telephone call with client while Dockets.xlsm is open in Excel",
            "rawSeconds": 2880,
            "hours": 0.8,
            "useExactHours": True,
            "clientRate": 475.0,
            "sharePct": 100.0,
            "status": "Draft",
        }
    )
    entry_id = add_res["entryId"]
    assert entry_id

    # CSPM.xlsm save succeeded and pending queue was written
    sync_service = LegacyDocketsSyncService(paths)
    pending_file = sync_service._pending_queue_path()
    assert pending_file.exists()

    # Remove lock file and flush pending syncs
    lock_file.unlink()
    flush_res = sync_service.flush_pending_syncs()
    assert len(flush_res) == 1
    assert flush_res[0]["ok"] is True

    wb = openpyxl.load_workbook(legacy_path, keep_vba=True, data_only=False)
    ws = wb[LEGACY_DOCKETS_SHEET]
    assert ws.tables[LEGACY_DOCKETS_TABLE].ref == "A1:M3"
    assert ws.cell(row=3, column=4).value == "Telephone call with client while Dockets.xlsm is open in Excel"
    assert int(ws.cell(row=3, column=12).value) == 2880
    LegacyDocketsSyncService._close_workbook(wb)


def test_start_new_timer_switches_matter_and_preserves_custom_ui_ribbon(tmp_path: Path) -> None:
    import zipfile

    paths, repo, legacy_path = _create_isolated_cspm_repo(tmp_path)

    # Inject customUI ribbon, ui/extensibility relationship, queryTables, and ExternalData_1 definedName
    # into legacy_path to simulate real Dockets.xlsm structure.
    patched_zip = legacy_path.with_name("patched_dockets.xlsm")
    with zipfile.ZipFile(legacy_path, "r") as zin, zipfile.ZipFile(
        patched_zip, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "_rels/.rels":
                txt = data.decode("utf-8").replace(
                    "</Relationships>",
                    '<Relationship Id="R9ba3f36e5ed24550" Type="http://schemas.microsoft.com/office/2007/relationships/ui/extensibility" Target="customUI/customUI14.xml"/></Relationships>',
                )
                data = txt.encode("utf-8")
            elif name == "xl/workbook.xml":
                txt = data.decode("utf-8")
                if "<definedNames>" not in txt:
                    txt = txt.replace(
                        "</workbook>",
                        '<definedNames><definedName name="ExternalData_1" localSheetId="0" hidden="1">Dockets!$A$1:$H$2</definedName></definedNames></workbook>',
                    )
                data = txt.encode("utf-8")
            zout.writestr(name, data)
        zout.writestr("customUI/customUI14.xml", b'<customUI xmlns="http://schemas.microsoft.com/office/2009/07/customui"><ribbon/></customUI>')
        zout.writestr("xl/queryTables/queryTable1.xml", b'<queryTable xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" name="ExternalData_1"/>')
        zout.writestr("xl/connections.xml", b'<connections xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>')
    patched_zip.replace(legacy_path)

    # Verify _get_or_create_matter matches 'MEGA-COR-26-0093 — Corporate Reorganization' without creating a duplicate
    matched_matter = repo._get_or_create_matter(
        matter_name="MEGA-COR-26-0093 \u2014 Corporate Reorganization",
        client_id="CLI-001",
        parent_id="PAR-001",
        default_rate=475.0,
        default_share_pct=100.0,
    )
    assert matched_matter[sc.COL_MATTER_ID] == "MAT-001"
    assert matched_matter[sc.COL_MATTER_NUMBER] == "MEGA-COR-26-0093"

    # Sync a new docket via repo.add_time_entry using the display-formatted matterName
    add_res = repo.add_time_entry(
        {
            "date": "2026-10-18",
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterName": "MEGA-COR-26-0093 \u2014 Corporate Reorganization",
            "description": "Surgical XML sync preserving customUI ribbon",
            "rawSeconds": 3600,
            "hours": 1.0,
            "useExactHours": True,
            "clientRate": 475.0,
            "sharePct": 100.0,
            "status": "Draft",
        }
    )
    assert add_res["ok"] is True

    with zipfile.ZipFile(legacy_path, "r") as zcheck:
        names = set(zcheck.namelist())
        assert "customUI/customUI14.xml" in names
        assert "xl/queryTables/queryTable1.xml" in names
        assert "xl/connections.xml" in names
        rels_xml = zcheck.read("_rels/.rels").decode("utf-8")
        assert "customUI/customUI14.xml" in rels_xml
        wb_xml = zcheck.read("xl/workbook.xml").decode("utf-8")
        assert "Dockets!$A$1:$H$3" in wb_xml

    # Verify TrayController.startNewTimerWithDescription switches to a new 00:00:00 timer
    _app = QCoreApplication.instance() or QCoreApplication([])
    mock_icon = MagicMock()
    mock_app_controller = SimpleNamespace(_excel_repo=repo, _theme_name="Dark", _tray_icon=mock_icon)
    tray = TrayController(mock_app_controller)
    commands: list[dict] = []
    tray.mainTimerCommandRequested.connect(lambda cmd: commands.append(dict(cmd)))

    tray.startNewTimerWithDescription(
        "Megatech Integrated Services Ltd.",
        "MEGA-COR-26-0093 \u2014 Corporate Reorganization",
        "New matter timer",
    )
    assert tray.hasRunningTimer is True
    assert tray.displayClientName == "Megatech Integrated Services Ltd."
    assert tray.displayMatterName == "MEGA-COR-26-0093 \u2014 Corporate Reorganization"
    assert tray.displayDescription == "New matter timer"
    assert commands[-1]["action"] == "start_new"
    tray.shutdown()


def test_tray_manual_time_entry_and_matter_deduplication(tmp_path: Path) -> None:
    paths, repo, _legacy_path = _create_isolated_cspm_repo(tmp_path)
    _app = QCoreApplication.instance() or QCoreApplication([])
    mock_icon = MagicMock()
    mock_app_controller = SimpleNamespace(
        _excel_repo=repo,
        _theme_name="Dark",
        _tray_icon=mock_icon,
        saveTimeDocketEntry=lambda payload: repo.add_time_entry(payload),
    )
    tray = TrayController(mock_app_controller)

    # 1. Verify parsing of various manual time formats
    assert tray.parseManualTimeSeconds("2.8") == 10080
    assert tray.parseManualTimeSeconds("1.5h") == 5400
    assert tray.parseManualTimeSeconds("45m") == 2700
    assert tray.parseManualTimeSeconds("1h 30m") == 5400
    assert tray.parseManualTimeSeconds("02:48:00") == 10080
    assert tray.parseManualTimeSeconds("0:30") == 1800

    # 2. Verify reversed WIP billing wizard label ("Corporate Reorganization - 26-0093" and em-dash variant)
    # resolves to existing MAT-001 without spawning a duplicate matter
    m_rev_dash = repo._get_or_create_matter(
        matter_name="Corporate Reorganization - 26-0093",
        client_id="CLI-001",
        parent_id="PAR-001",
        default_rate=475.0,
        default_share_pct=100.0,
    )
    assert m_rev_dash[sc.COL_MATTER_ID] == "MAT-001"
    m_rev_em = repo._get_or_create_matter(
        matter_name="Corporate Reorganization \u2014 26-0093",
        client_id="CLI-001",
        parent_id="PAR-001",
        default_rate=475.0,
        default_share_pct=100.0,
    )
    assert m_rev_em[sc.COL_MATTER_ID] == "MAT-001"

    # 3. Direct manual time entry save from tray (+ New -> Save)
    saved_ok = tray.saveManualTimeEntry(
        "Megatech Integrated Services Ltd.",
        "Corporate Reorganization \u2014 26-0093",
        "Manual time entry from tray popup",
        "2.8",
    )
    assert saved_ok is True
    assert tray.hasRunningTimer is False
    assert tray.displayStatus == "LATEST"
    assert tray.displayElapsedSeconds == 10080
    assert tray.displayElapsedFormatted == "02:48:00"
    assert tray.displayMatterName == "MEGA-COR-26-0093 \u2014 Corporate Reorganization"

    # Verify only 1 matter exists in TBL_MATTERS (no phantom matter created)
    all_matters = repo._read_table_rows(TBL_MATTERS)
    assert len(all_matters) == 1
    assert all_matters[0][sc.COL_MATTER_ID] == "MAT-001"

    # 4. Start timer and override elapsed time inline via setPrimaryManualTime
    tray.startNewTimerWithManualTime(
        "Megatech Integrated Services Ltd.",
        "MEGA-COR-26-0093 \u2014 Corporate Reorganization",
        "Timer started with manual offset",
        "45m",
    )
    assert tray.hasRunningTimer is True
    assert tray.displayElapsedSeconds >= 2700
    assert tray.setPrimaryManualTime("1.5h") is True
    assert 5400 <= tray.displayElapsedSeconds <= 5402
    tray.shutdown()


def test_multi_tab_today_dockets_and_simultaneous_timer_warning(tmp_path: Path) -> None:
    paths, repo, _legacy_path = _create_isolated_cspm_repo(tmp_path)
    _app = QCoreApplication.instance() or QCoreApplication([])
    mock_icon = MagicMock()
    mock_app_controller = SimpleNamespace(
        _excel_repo=repo,
        _paths=paths,
        _theme_name="Dark",
        _tray_icon=mock_icon,
        saveTimeDocketEntry=lambda payload: repo.add_time_entry(payload),
    )
    tray = TrayController(mock_app_controller)

    # 1. Start Timer A on Matter 1
    tray.startNewTimerWithDescription(
        "Megatech Integrated Services Ltd.",
        "MEGA-COR-26-0093 \u2014 Corporate Reorganization",
        "Drafting share rollovers",
    )
    assert tray.hasRunningTimer is True
    assert len(tray.todayDocketTabs) >= 1

    # 2. Click '+' to open a new blank docket tab for today
    new_tab_id = tray.createNewBlankDocketTab()
    assert new_tab_id.startswith("draft:")
    assert tray.selectedTabId == new_tab_id
    assert any(t["tabId"] == new_tab_id for t in tray.todayDocketTabs)

    # 3. Attempting to start Timer B while Timer A is RUNNING triggers Simultaneous Timer Warning
    started_immediately = tray.startNewTimerWithManualTime(
        "Megatech Integrated Services Ltd.",
        "General Corporate Advisory",
        "Second concurrent matter",
        "15m",
    )
    assert started_immediately is False
    assert tray.simultaneousTimerWarningVisible is True
    assert "Another timer is already running for" in tray.simultaneousTimerWarningMessage

    # 3a. Cancel leaves Timer A running and does not start Timer B
    tray.resolveSimultaneousTimerCancel()
    assert tray.simultaneousTimerWarningVisible is False
    running_timers = [t for t in tray._timers if t.isRunning]
    assert len(running_timers) == 1

    # 3b. Trigger warning again and choose 'Pause Other & Start This One'
    tray.selectDocketTab(new_tab_id)
    tray.startNewTimerWithManualTime(
        "Megatech Integrated Services Ltd.",
        "General Corporate Advisory",
        "Second matter after pausing first",
        "15m",
    )
    assert tray.simultaneousTimerWarningVisible is True
    tray.resolveSimultaneousTimerPauseOther()
    assert tray.simultaneousTimerWarningVisible is False
    running_timers = [t for t in tray._timers if t.isRunning]
    paused_timers = [t for t in tray._timers if not t.isRunning]
    assert len(running_timers) == 1
    assert len(paused_timers) == 1
    assert running_timers[0].description == "Second matter after pausing first"

    # 3c. Resume paused Timer A -> triggers warning -> choose 'Run Both Simultaneously'
    tray.resume_timer(paused_timers[0].timer_id)
    assert tray.simultaneousTimerWarningVisible is True
    tray.resolveSimultaneousTimerRunBoth()
    assert tray.simultaneousTimerWarningVisible is False
    running_timers = [t for t in tray._timers if t.isRunning]
    assert len(running_timers) == 2

    # 4. Verify multi-monitor geometry persistence & screen intersection check
    assert tray.saveMiniPopupGeometry(2200, 140, 420, 310) is True
    assert TrayController._rect_intersects_any_screen(
        2200, 140, 420, 310, [(0, 0, 1920, 1080), (1920, 0, 2560, 1440)]
    ) is True
    assert TrayController._rect_intersects_any_screen(
        9000, 9000, 420, 310, [(0, 0, 1920, 1080), (1920, 0, 2560, 1440)]
    ) is False
    tray.shutdown()


def test_legacy_dockets_conflict_detection_and_four_resolutions(tmp_path: Path) -> None:
    paths, repo, legacy_path = _create_isolated_cspm_repo(tmp_path)

    # Case A: First entry for (2026-10-19, MEGA-COR-26-0093) appends quietly without conflict
    res1 = repo.add_time_entry(
        {
            "date": "2026-10-19",
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093",
            "parentId": "PAR-001",
            "description": "Initial morning drafting",
            "rawSeconds": 3600,
            "hours": 1.0,
            "useExactHours": True,
            "clientRate": 475.0,
            "sharePct": 100.0,
            "status": "Draft",
        }
    )
    assert res1["ok"] is True
    sync1 = res1["legacyDocketsSync"]
    assert sync1["conflict"] is False
    assert sync1["status"] in ("APPENDED", "SYNCED")
    assert sync1["toastMessage"] == "Saved & synced to Dockets.xlsm"

    # Case B: Second separate entry on the SAME (Date, Matter_ID) triggers CONFLICT
    res2 = repo.add_time_entry(
        {
            "date": "2026-10-19",
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093",
            "parentId": "PAR-001",
            "description": "Afternoon follow-up call",
            "rawSeconds": 1800,
            "hours": 0.5,
            "useExactHours": True,
            "clientRate": 475.0,
            "sharePct": 100.0,
            "status": "Draft",
            "forceNewEntry": True,
        }
    )
    sync2 = res2["legacyDocketsSync"]
    assert sync2["conflict"] is True
    assert sync2["status"] == "CONFLICT"
    conflict_id = sync2["conflictId"]
    target_row = sync2["selectedTargetRow"]
    assert target_row == 3

    # Resolution 3: Combine / Add Time into Existing Row (1.0h + 0.5h = 1.5h, 3600s + 1800s = 5400s)
    comb_res = repo.resolve_legacy_dockets_conflict(conflict_id, "combine", target_row=target_row)
    assert comb_res["ok"] is True
    assert comb_res["status"] == "SYNCED"
    assert comb_res.get("resolutionMode") == "combine"

    wb = openpyxl.load_workbook(legacy_path, keep_vba=True, data_only=False)
    ws = wb[LEGACY_DOCKETS_SHEET]
    assert ws.max_row == 3
    assert float(ws.cell(row=3, column=5).value) == 1.5
    assert int(ws.cell(row=3, column=12).value) == 5400
    assert "Initial morning drafting; Afternoon follow-up call" in str(ws.cell(row=3, column=4).value)
    # Mark row 3 as Billed in Excel to test billed overwrite protection
    ws.cell(row=3, column=10).value = "26-0099"
    wb.save(legacy_path)
    LegacyDocketsSyncService._close_workbook(wb)

    # Third entry on same (Date, Matter_ID) when existing row is Billed
    res3 = repo.add_time_entry(
        {
            "date": "2026-10-19",
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093",
            "parentId": "PAR-001",
            "description": "Evening revisions",
            "rawSeconds": 1200,
            "hours": 0.4,
            "useExactHours": True,
            "clientRate": 475.0,
            "sharePct": 100.0,
            "status": "Draft",
            "forceNewEntry": True,
        }
    )
    sync3 = res3["legacyDocketsSync"]
    assert sync3["conflict"] is True
    assert any(r.get("isBilled") for r in sync3["existingRows"])

    # Overwrite without confirm_billed_overwrite=True requires confirmation first
    ow_unconfirmed = repo.resolve_legacy_dockets_conflict(
        sync3, "overwrite", target_row=3, confirm_billed_overwrite=False
    )
    assert ow_unconfirmed.get("requiresBilledConfirmation") is True

    # Resolution 2: Append as Separate New Row
    app_res = repo.resolve_legacy_dockets_conflict(sync3, "append", target_row=3)
    assert app_res["ok"] is True
    assert app_res["status"] in ("APPENDED", "SYNCED")
    assert app_res.get("resolutionMode") == "append"

    wb2 = openpyxl.load_workbook(legacy_path, keep_vba=True, data_only=False)
    ws2 = wb2[LEGACY_DOCKETS_SHEET]
    assert ws2.max_row == 4
    assert ws2.tables[LEGACY_DOCKETS_TABLE].ref == "A1:M4"
    LegacyDocketsSyncService._close_workbook(wb2)

    # Resolution 4: Skip Dockets.xlsm (Save in CSPM Only)
    skip_res = repo.resolve_legacy_dockets_conflict(sync3, "skip", target_row=3)
    assert skip_res["ok"] is True
    assert skip_res["status"] == "SKIPPED"

    # Resolution 1: Overwrite Existing Row (with confirm_billed_overwrite=True)
    ow_confirmed = repo.resolve_legacy_dockets_conflict(
        sync3, "overwrite", target_row=3, confirm_billed_overwrite=True
    )
    assert ow_confirmed["ok"] is True
    assert ow_confirmed["status"] in ("OVERWRITTEN", "SYNCED")



def test_sqlite_realtime_autosave_crash_recovery_and_close_safeguard(tmp_path: Path) -> None:
    paths, repo, legacy_path = _create_isolated_cspm_repo(tmp_path)
    sqlite_repo = repo._get_sqlite_repo()
    assert sqlite_repo is not None

    # 1. Simulate a timer that was RUNNING and autosaved to SQLite 5s before a hard PC crash
    started_utc = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(seconds=600)).strftime("%Y-%m-%dT%H:%M:%SZ")
    heartbeat_utc = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(seconds=300)).strftime("%Y-%m-%dT%H:%M:%SZ")
    crashed_row = repo.upsert_sqlite_docket_heartbeat(
        {
            "entryId": "TM-CRASH-01",
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093",
            "date": "2026-10-19",
            "description": "Crashed mid-timer research",
            "rawSeconds": 300,
            "entryStatus": "RUNNING",
            "startedAtUtc": started_utc,
            "lastHeartbeatUtc": heartbeat_utc,
            "legacyDocketsSyncStatus": "PENDING_CRASH_RECOVERY",
        }
    )
    assert crashed_row["entryId"] == "TM-CRASH-01"

    # Also seed a PENDING_SYNC row (saved in SQLite before crash, not yet in Dockets.xlsm)
    pending_row = repo.upsert_sqlite_docket_heartbeat(
        {
            "entryId": "TM-PEND-01",
            "clientId": "CLI-001",
            "clientName": "Megatech Integrated Services Ltd.",
            "matterId": "MAT-001",
            "matterName": "MEGA-COR-26-0093",
            "date": "2026-10-17",
            "description": "Saved before crash but pending Excel sync",
            "rawSeconds": 1800,
            "entryStatus": "SAVED",
            "legacyDocketsSyncStatus": "PENDING_SYNC",
        }
    )
    assert pending_row["entryId"] == "TM-PEND-01"

    # 2. Run startup reconciliation:
    # - TM-PEND-01 (2026-10-17) quietly appends to Dockets.xlsm Table1_1 and becomes SYNCED
    # - TM-CRASH-01 (2026-10-19) is restored as PAUSED in SQLite and returned in crashRecoveryDockets
    rec = repo.reconcile_unsynced_dockets_on_startup()
    assert rec["ok"] is True
    assert any(item["entryId"] == "TM-PEND-01" for item in rec["syncedEntries"])
    assert any(item["entryId"] == "TM-CRASH-01" for item in rec["crashRecoveryDockets"])

    crash_info = next(item for item in rec["crashRecoveryDockets"] if item["entryId"] == "TM-CRASH-01")
    assert crash_info["wasRunning"] is True
    assert crash_info["heartbeatSeconds"] == 300
    assert crash_info["wallClockSeconds"] >= 590

    # Verify TM-CRASH-01 was safely transitioned to PAUSED in SQLite
    db_row = sqlite_repo.get_time_entry_by_id("TM-CRASH-01")
    assert db_row is not None
    assert db_row["EntryStatus"] == "PAUSED"

    # 3. Resolve TM-CRASH-01 using Option 2 ('sync_and_keep_paused') with 'heartbeat' time
    res_keep = repo.resolve_crash_recovery_docket(
        "TM-CRASH-01", "sync_and_keep_paused", time_option="heartbeat"
    )
    assert res_keep["ok"] is True
    assert res_keep["keepPausedTab"] is True
    assert res_keep["legacyDocketsSync"]["conflict"] is False

    db_row_after = sqlite_repo.get_time_entry_by_id("TM-CRASH-01")
    assert db_row_after is not None
    assert db_row_after["EntryStatus"] == "PAUSED"
    assert db_row_after["LegacyDocketsSyncStatus"] == "SYNCED"
    assert db_row_after["LegacyDocketsRowSignature"] is not None

    # 4. Verify Close Safeguard in TrayController when a timer is RUNNING
    _app = QCoreApplication.instance() or QCoreApplication([])
    mock_icon = MagicMock()
    mock_app_controller = SimpleNamespace(
        _excel_repo=repo,
        _paths=paths,
        _theme_name="Dark",
        _tray_icon=mock_icon,
        saveTimeDocketEntry=lambda payload: repo.add_time_entry(payload),
    )
    tray = TrayController(mock_app_controller)
    tray.startNewTimerWithDescription(
        "Megatech Integrated Services Ltd.",
        "MEGA-COR-26-0093 \u2014 Corporate Reorganization",
        "Active timer during close attempt",
    )
    assert tray.hasUnsavedOrActiveDocketOnClose() is True
    assert tray.requestCloseWithSafeguard("main_window") is True
    assert tray.closeSafeguardVisible is True

    # Option 4: Cancel close safeguard
    cancel_res = tray.resolveCloseSafeguard("cancel")
    assert cancel_res["action"] == "cancel"
    assert tray.closeSafeguardVisible is False
    assert tray.hasRunningTimer is True

    # Option 2: Pause, Sync Current Time to Dockets.xlsm & Quit (Resume Next Launch)
    tray.setPrimaryManualTime("30m")
    pause_quit_res = tray.resolveCloseSafeguard("pause_sync_and_quit")
    assert pause_quit_res["ok"] is True
    assert tray.hasRunningTimer is False
    assert tray._pause_sync_quit_completed is True
    tray.shutdown()



